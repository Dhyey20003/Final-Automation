from __future__ import annotations

import io

from openpyxl import Workbook

from conftest import HEAD, csv_bytes, emp, run_import


def roster_rows(n=5, loc="Ahmedabad", start=1):
    return [[f"RHC{i:04d}", f"Person {chr(64 + i)} Test", "Associate", f"p{i}@relay-sample.test", loc, "01-02-2024", "Active"]
            for i in range(start, start + n)]


def test_roster_import_creates_and_preserves_leading_zeros(api):
    rows = roster_rows(3) + [["0099", "Zero Lead", "Associate", "z@relay-sample.test", "Vadodara", "2024-03-05", "Active"]]
    bid, plan, res = run_import(api, csv_bytes([HEAD] + rows), "roster", options={"scope": ["Ahmedabad", "Vadodara"]})
    assert plan["summary"]["new"] == 4 and plan["blocking"] == 0
    assert res["summary"]["created"] == 4
    assert api.get("/api/employees?q=0099")["rows"][0]["rhc_id"] == "0099"


def test_xlsx_leading_zero_number_format_and_formulas_not_run(api):
    wb = Workbook()
    ws = wb.active
    ws.append(["RHC ID", "Employee Name", "Location", "Date of Joining", "Designation"])
    ws.append([123, "Numeric Zero", "Ahmedabad", None, "=1+1"])
    ws["A2"].number_format = "00000"
    import datetime
    ws["D2"] = datetime.datetime(2024, 6, 1)
    buf = io.BytesIO()
    wb.save(buf)
    bid, plan, res = run_import(api, buf.getvalue(), "roster", filename="f.xlsx", options={"scope": ["Ahmedabad"]})
    row = plan["rows"][0]
    assert row["rhc_id"] == "00123"
    assert any("formula" in w for w in row["warnings"])
    e = api.get("/api/employees?q=00123")["rows"][0]
    assert e["designation"] is None and e["date_of_joining"] == "2024-06-01"


def test_rejects_macro_and_old_formats(api):
    api.post("/api/imports/upload", files={"file": ("x.xlsm", b"PK")}, status=400)
    api.post("/api/imports/upload", files={"file": ("x.xls", b"xx")}, status=400)


def test_duplicates_and_invalid_rows_block_commit(api):
    rows = [["RHC0001", "A One", "", "", "Ahmedabad", "", ""], ["RHC0001", "A Two", "", "", "Ahmedabad", "", ""],
            ["RHC0003", "Bad Date", "", "", "Ahmedabad", "31-31-2024", ""], ["RHC0004", "Bad Loc", "", "", "Surat", "", ""]]
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"create_unmatched": True}, commit=False)
    statuses = [r["status"] for r in plan["rows"]]
    assert statuses.count("duplicate") == 2 and statuses.count("error") == 2
    assert plan["blocking"] == 4
    api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan["plan_hash"]}, status=422)
    # skipping one duplicate resolves the other
    body = {"purpose": "partial", "header_row": 1, "mapping": api.get(f"/api/imports/{bid}/sheet")["suggested_mapping"],
            "options": {"create_unmatched": True}, "resolutions": {"3": "skip", "4": "skip", "5": "skip"}}
    plan2 = api.post(f"/api/imports/{bid}/preview", body)
    assert plan2["blocking"] == 0 and plan2["summary"]["new"] == 1


def test_ambiguous_email_and_name_matches_are_never_merged(api):
    emp(api, "RHC0001", "Asha Patel", official_email="asha@relay-sample.test")
    rows = [["RHC0999", "Someone", "", "asha@relay-sample.test", "Ahmedabad", "", ""],  # email belongs to another ID
            ["", "Asha Patel", "Manager", "", "Ahmedabad", "", ""]]  # name only
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"email_fallback": True}, commit=False)
    r1, r2 = plan["rows"]
    assert r1["status"] == "error" and "never merged" in r1["errors"][0]
    assert r2["status"] in ("error", "unmatched")
    assert r2["suggestions"] and r2["suggestions"][0]["rhc_id"] == "RHC0001"
    assert not r2.get("employee_id")  # suggestion only


def test_name_suggestion_can_be_linked_explicitly(api):
    e = emp(api, "RHC0001", "Asha Patel")
    rows = [["RHC0001X", "Asha Patel", "Manager", "", "Ahmedabad", "", ""]]
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "designation", commit=False)
    assert plan["rows"][0]["status"] == "unmatched"
    body = {"purpose": "designation", "header_row": 1, "mapping": api.get(f"/api/imports/{bid}/sheet")["suggested_mapping"],
            "resolutions": {"2": f"link:{e['id']}"}}
    plan2 = api.post(f"/api/imports/{bid}/preview", body)
    assert plan2["rows"][0]["status"] == "update"
    api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan2["plan_hash"]})
    assert api.get(f"/api/employees/{e['id']}")["employee"]["designation"] == "Manager"


def test_email_fallback_matches_unique_email(api):
    e = emp(api, "RHC0050", "Email Match", official_email="em@relay-sample.test")
    rows = [["", "Email Match", "Lead", "em@relay-sample.test", "", "", ""]]
    mapping = {"0": "rhc_id", "1": "full_name", "2": "designation", "3": "official_email"}
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"email_fallback": True}, mapping=mapping)
    assert plan["rows"][0]["match_method"] == "email"
    assert api.get(f"/api/employees/{e['id']}")["employee"]["designation"] == "Lead"


def test_partial_update_blank_cells_and_no_deactivation(api):
    a = emp(api, "RHC0001", "Keep Me", designation="Associate", official_email="k@relay-sample.test")
    b = emp(api, "RHC0002", "Not In File")
    rows = [["RHC0001", "", "", "", "", "", ""]]
    bid, plan, res = run_import(api, csv_bytes([HEAD] + rows), "partial")
    assert plan["rows"][0]["status"] == "unchanged"
    got = api.get(f"/api/employees/{a['id']}")["employee"]
    assert got["designation"] == "Associate" and got["official_email"] == "k@relay-sample.test"
    assert api.get(f"/api/employees/{b['id']}")["employee"]["employment_status"] == "Active"
    # explicitly selected: blank clears designation
    bid, plan, res = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"blank_clears": ["designation"]},
                                commit_body={"acknowledge_repeat": True})
    assert api.get(f"/api/employees/{a['id']}")["employee"]["designation"] is None
    # identity fields can never be cleared
    _, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"blank_clears": ["full_name"]}, commit=False)
    assert plan["mapping_errors"]


def test_designation_and_selected_field_updates_touch_only_allowed_fields(api):
    a = emp(api, "RHC0001", "Only Designation", department="Ops")
    rows = [["RHC ID", "Designation", "Department"], ["RHC0001", "Lead", "Finance"]]
    bid, plan, _ = run_import(api, csv_bytes(rows), "designation")
    got = api.get(f"/api/employees/{a['id']}")["employee"]
    assert got["designation"] == "Lead" and got["department"] == "Ops"
    bid, plan, _ = run_import(api, csv_bytes(rows), "selected", options={"selected_fields": ["department"]},
                              commit_body={"acknowledge_repeat": True})
    got = api.get(f"/api/employees/{a['id']}")["employee"]
    assert got["department"] == "Finance"


def test_historical_then_roster_workflow(api):
    hist = [["RHC0001", "Current Person", "Trainee", "cur@relay-sample.test", "Ahmedabad", "01-01-2020", "Active"],
            ["RHC0900", "Former Person", "Associate", "old@relay-sample.test", "Vadodara", "01-01-2018", "Left"]]
    hist_head = HEAD + ["Date of Birth", "Remarks"]
    hist = [r + ["10-10-1990", "old remark"] for r in hist]
    _, plan, res = run_import(api, csv_bytes([hist_head] + hist), "historical")
    assert plan["summary"]["historical_new"] == 2
    assert api.get("/api/employees")["total"] == 0  # not in the active directory
    assert api.get("/api/employees?view=historical")["total"] == 2
    roster = [["RHC0001", "Current Person", "Senior Associate", "", "Ahmedabad", "", ""],
              ["RHC0002", "New Joiner", "Associate", "nj@relay-sample.test", "Ahmedabad", "01-09-2026", ""]]
    _, plan, res = run_import(api, csv_bytes([HEAD] + roster), "roster",
                              options={"scope": ["Ahmedabad"], "enrich_fields": ["date_of_joining", "date_of_birth", "official_email"]})
    assert plan["summary"]["promote"] == 1 and plan["summary"]["new"] == 1
    active = {r["rhc_id"]: r for r in api.get("/api/employees")["rows"]}
    cur = active["RHC0001"]
    assert cur["designation"] == "Senior Associate"
    assert cur["date_of_birth"] == "1990-10-10" and cur["official_email"] == "cur@relay-sample.test"  # approved history kept
    assert cur["notes"] is None  # not approved: cleared
    assert "RHC0900" not in active
    # history imported after the roster only fills blanks on current employees
    later = [["RHC0002", "New Joiner", "Trainee", "", "Ahmedabad", "", "", "05-05-1995", ""]]
    _, plan, _ = run_import(api, csv_bytes([hist_head] + later), "historical")
    r = plan["rows"][0]
    assert r["status"] == "enrich" and [c["field"] for c in r["changes"]] == ["date_of_birth"]
    nj = api.get("/api/employees?q=RHC0002")["rows"][0]
    assert nj["designation"] == "Associate" and nj["date_of_birth"] == "1995-05-05"


def test_complete_roster_requires_confirmation_to_mark_left(api):
    for i in range(1, 5):
        emp(api, f"RHC000{i}", f"Emp {i}", "Ahmedabad")
    emp(api, "RHC0100", "Vadodara Emp", "Vadodara")
    rows = roster_rows(2)
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "roster", options={"scope": ["Ahmedabad"]}, commit=False)
    missing = {d["rhc_id"] for d in plan["deactivations"]}
    assert missing == {"RHC0003", "RHC0004"}  # Vadodara is outside the declared scope
    ids = [d["employee_id"] for d in plan["deactivations"]]
    api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan["plan_hash"], "deactivate_ids": ids}, status=422)  # no confirmation
    api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan["plan_hash"], "deactivate_ids": ids[:1], "confirm_deactivations": True,
                                            "deactivation_last_working_date": "2026-09-30"})
    left = api.get("/api/employees?view=left")["rows"]
    assert [r["rhc_id"] for r in left] == ["RHC0003"] and left[0]["last_working_date"] == "2026-09-30"
    assert api.get("/api/employees?q=RHC0004")["rows"][0]["employment_status"] == "Active"


def test_roster_needs_declared_scope(api):
    _, plan, _ = run_import(api, csv_bytes([HEAD] + roster_rows(1)), "roster", commit=False)
    assert any("locations" in m for m in plan["mapping_errors"])


def test_repeated_import_detected_and_plan_changes_detected(api):
    content = csv_bytes([HEAD] + roster_rows(2))
    run_import(api, content, "partial", options={"create_unmatched": True})
    bid, plan, _ = run_import(api, content, "partial", options={"create_unmatched": True}, commit=False)
    assert plan["previous_imports"]
    r = api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan["plan_hash"]}, status=422)
    assert r["error"]["code"] == "repeat_import"
    # data changes between preview and commit -> refuse and re-preview
    bid, plan, _ = run_import(api, csv_bytes([HEAD, ["RHC0001", "Person A Test", "Changed", "", "", "", ""]]), "partial", commit=False)
    e = api.get("/api/employees?q=RHC0001")["rows"][0]
    api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"department": "Moved"}})
    api.post(f"/api/imports/{bid}/commit", {"plan_hash": plan["plan_hash"]}, status=409)


def test_import_audit_and_error_report(api, ctx):
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + roster_rows(2) + [["", "No Id", "", "", "", "", ""]]), "partial",
                              options={"create_unmatched": True}, resolutions={"4": "skip"})
    with ctx.db() as conn:
        b = conn.execute("SELECT * FROM import_batches WHERE id = ?", (bid,)).fetchone()
        a = conn.execute("SELECT details FROM audit_log WHERE action = 'import.commit'").fetchone()
    assert b["file_hash"] and b["operator"] == "admin@relay.test" and b["purpose"] == "partial"
    assert "sha256" in a["details"]
    r = api.get(f"/api/imports/{bid}/report")
    assert r.content[:2] == b"PK"


def test_undo_import_and_conflicts(api):
    base = emp(api, "RHC0001", "Person A Test", designation="Old")
    rows = [["RHC0001", "Person A Test", "New", "", "", "", ""], ["RHC0002", "Person B Test", "Assoc", "", "Ahmedabad", "", ""]]
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "partial", options={"create_unmatched": True})
    assert api.get("/api/employees")["total"] == 2
    up = api.get(f"/api/imports/{bid}/undo-preview")
    assert len(up["revertible"]) == 2 and not up["conflicts"]
    # a later manual change creates a conflict that undo refuses to overwrite
    e = api.get(f"/api/employees/{base['id']}")["employee"]
    api.patch(f"/api/employees/{base['id']}", {"version": e["version"], "changes": {"designation": "Manual"}})
    up = api.get(f"/api/imports/{bid}/undo-preview")
    assert len(up["conflicts"]) == 1
    api.post(f"/api/imports/{bid}/undo", {"mode": "all"}, status=409)
    res = api.post(f"/api/imports/{bid}/undo", {"mode": "non_conflicting"})
    assert res["status"] == "partially_undone"
    assert api.get(f"/api/employees/{base['id']}")["employee"]["designation"] == "Manual"
    assert api.get("/api/employees?q=RHC0002")["total"] == 0  # created record removed


def test_full_undo_restores_deactivated_and_promoted(api):
    emp(api, "RHC0001", "Stay")
    gone = emp(api, "RHC0002", "Gone")
    run_import(api, csv_bytes([HEAD, ["RHC0900", "Hist Person", "", "", "Ahmedabad", "", "Left"]]), "historical")
    rows = [["RHC0001", "Stay", "", "", "Ahmedabad", "", ""], ["RHC0900", "Hist Person", "", "", "Ahmedabad", "", ""]]
    bid, plan, _ = run_import(api, csv_bytes([HEAD] + rows), "roster", options={"scope": ["Ahmedabad"]},
                              commit_body={"deactivate_ids": [gone["id"]], "confirm_deactivations": True})
    assert api.get("/api/employees")["total"] == 2
    res = api.post(f"/api/imports/{bid}/undo", {"mode": "all"})
    assert res["status"] == "undone"
    active = {r["rhc_id"] for r in api.get("/api/employees")["rows"]}
    assert active == {"RHC0001", "RHC0002"}
    assert api.get("/api/employees?view=historical")["rows"][0]["rhc_id"] == "RHC0900"


def test_mapping_templates_are_reusable(api):
    content = csv_bytes([["Emp Code", "Name", "Desig"], ["RHC0001", "Tpl Person", "Analyst"]])
    up = api.post("/api/imports/upload", files={"file": ("t.csv", content)})
    mapping = {"0": "rhc_id", "1": "full_name", "2": "designation"}
    plan = api.post(f"/api/imports/{up['batch_id']}/preview", {"purpose": "partial", "header_row": 1, "mapping": mapping,
                                                             "options": {"create_unmatched": True}, "save_template_as": "HR monthly"})
    tid = plan["template_id"]
    up2 = api.post("/api/imports/upload", files={"file": ("t2.csv", csv_bytes([["Desig", "Emp Code", "Name"], ["X", "Y", "Z"]]))})
    info = api.get(f"/api/imports/{up2['batch_id']}/sheet?template_id={tid}")
    assert info["suggested_mapping"] == {"0": "designation", "1": "rhc_id", "2": "full_name"}
