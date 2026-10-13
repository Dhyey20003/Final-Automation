from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import load_workbook

from conftest import emp
from hub.services.employees import bench_days


def test_create_edit_and_leading_zero_ids(api):
    e = emp(api, "0457", "Zero Lead", date_of_joining="15-08-2024")
    assert e["rhc_id"] == "0457"
    assert e["date_of_joining"] == "2024-08-15"  # day-first parsing
    e2 = emp(api, "rhc 0012", "Asha Patel", "vdr")
    assert e2["rhc_id"] == "RHC0012" and e2["location"] == "Vadodara"
    # duplicate RHC ID (case-insensitive) is rejected
    r = api.post("/api/employees", {"values": {"rhc_id": "rhc0012", "full_name": "X"}}, status=422)
    assert "already used" in str(r)
    # inline edit with version
    res = api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"designation": "Analyst"}})
    assert res["employee"]["designation"] == "Analyst" and res["employee"]["version"] == 2
    lst = api.get("/api/employees?q=0457")
    assert lst["total"] == 1 and lst["rows"][0]["rhc_id"] == "0457"


def test_version_conflict_prevents_silent_overwrite(api):
    e = emp(api, "RHC1", "One")
    api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"designation": "A"}})
    r = api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"designation": "B"}}, status=409)
    assert r["error"]["details"]["current"]["designation"] == "A"
    assert api.get(f"/api/employees/{e['id']}")["employee"]["designation"] == "A"


def test_date_validation_and_statuses(api):
    r = api.post("/api/employees", {"values": {"rhc_id": "R2", "full_name": "Two", "date_of_joining": "2024-05-01",
                                                "last_working_date": "2024-04-01"}}, status=422)
    assert "last_working_date" in r["error"]["details"]["fields"]
    r = api.post("/api/employees", {"values": {"rhc_id": "R3", "full_name": "Three", "bench_start": "2026-02-10", "bench_end": "2026-02-01"}}, status=422)
    assert "bench_end" in r["error"]["details"]["fields"]
    r = api.post("/api/employees", {"values": {"rhc_id": "R4", "full_name": "Four", "location": "Mumbai"}}, status=422)
    e = emp(api, "R5", "Five", employment_status="Left", last_working_date="2025-01-01")
    assert e["employment_status"] == "Left"
    # employment and assignment status are separate
    assert e["assignment_status"] == "Unknown"


def test_bench_days_calculation():
    today = date(2026, 9, 29)
    rec = {"assignment_status": "Bench", "bench_start": "2026-09-29", "bench_end": None}
    assert bench_days(rec, today) == 0
    rec["bench_start"] = "2026-09-20"
    assert bench_days(rec, today) == 9
    rec["bench_end"] = "2026-09-25"
    assert bench_days(rec, today) == 5
    rec.update(bench_start="2026-10-05", bench_end=None)
    assert bench_days(rec, today) is None  # future start
    assert bench_days({"assignment_status": "Client", "bench_start": "2026-09-01"}, today) is None


def test_bench_days_filter_and_sort_in_sql(api):
    from datetime import datetime
    from zoneinfo import ZoneInfo

    today = datetime.now(ZoneInfo("Asia/Kolkata")).date()
    emp(api, "B1", "Bench Ten", assignment_status="Bench", bench_start=(today - timedelta(days=10)).isoformat())
    emp(api, "B2", "Bench Two", assignment_status="Bench", bench_start=(today - timedelta(days=2)).isoformat())
    emp(api, "C1", "Client One", assignment_status="Client", client_name="Contoso")
    res = api.get('/api/employees?sort=bench_days&dir=desc&filters={"bench_days": ">=5"}')
    assert [r["rhc_id"] for r in res["rows"]] == ["B1"] and res["rows"][0]["bench_days"] == 10


def test_assignment_history_is_preserved(api):
    e = emp(api, "H1", "Hist", assignment_status="Client", client_name="Contoso", assignment_start="2025-01-01")
    api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"assignment_status": "Bench", "bench_start": "2026-09-01",
                                                                                 "client_name": None, "assignment_end": "2026-08-31"}})
    d = api.get(f"/api/employees/{e['id']}")
    assert d["assignment_history"][0]["client_name"] == "Contoso"
    fields = {h["field"] for h in d["history"]}
    assert {"assignment_status", "bench_start", "client_name"} <= fields


def test_archive_keeps_history(api):
    e = emp(api, "A1", "Arch")
    api.post(f"/api/employees/{e['id']}/archive", {"version": e["version"]}, status=422)  # active cannot be archived
    e = api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"employment_status": "Left", "last_working_date": "2026-01-31"}})["employee"]
    api.post(f"/api/employees/{e['id']}/archive", {"version": e["version"]})
    assert api.get("/api/employees?view=left")["total"] == 0
    assert api.get("/api/employees?view=archived")["total"] == 1
    assert api.get(f"/api/employees/{e['id']}")["history"]


def test_custom_fields_everywhere(api):
    api.post("/api/fields", {"label": "Shift", "field_type": "dropdown", "options": ["Day", "Night"]})
    api.post("/api/fields", {"label": "Laptop issued", "field_type": "yesno"})
    api.post("/api/fields", {"label": "Experience years", "field_type": "number"})
    api.post("/api/fields", {"label": "Visa expiry", "field_type": "date"})
    api.post("/api/fields", {"label": "Badge", "field_type": "text"})
    meta = api.get("/api/meta")
    keys = {f["key"] for f in meta["fields"]}
    assert {"cf_shift", "cf_laptop_issued", "cf_experience_years", "cf_visa_expiry", "cf_badge"} <= keys
    e = emp(api, "CF1", "Custom", cf_shift="night", cf_laptop_issued="y", cf_experience_years="4.5", cf_visa_expiry="01/02/2027")
    assert e["cf_shift"] == "Night" and e["cf_laptop_issued"] == "Yes" and e["cf_experience_years"] == 4.5 and e["cf_visa_expiry"] == "2027-02-01"
    api.post("/api/employees", {"values": {"rhc_id": "CF2", "full_name": "Bad", "cf_shift": "Evening"}}, status=422)
    res = api.get('/api/employees?filters={"cf_shift": "Night"}')
    assert res["total"] == 1
    # export includes custom columns
    r = api.get("/api/employees/export?columns=rhc_id,full_name,cf_shift")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["RHC ID", "Full name", "Shift"]
    assert ws["C2"].value == "Night"


def test_export_protects_against_formula_injection(api):
    emp(api, "F1", "=HYPERLINK(\"http://x\",\"click\")", designation="+SUM(1,1)")
    r = api.get("/api/employees/export?columns=rhc_id,full_name,designation")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["B2"].value.startswith("'=") and ws["B2"].data_type == "s"
    assert ws["C2"].value == "'+SUM(1,1)"


def test_bulk_paste_is_previewed_and_atomic(api):
    a, b = emp(api, "P1", "Paste One"), emp(api, "P2", "Paste Two")
    upd = [{"id": a["id"], "version": a["version"], "changes": {"designation": "X"}},
           {"id": b["id"], "version": b["version"], "changes": {"date_of_joining": "not a date"}}]
    pv = api.post("/api/employees/bulk-preview", {"updates": upd})
    assert pv["ok"] is False and pv["results"][0]["ok"] and not pv["results"][1]["ok"]
    assert api.get(f"/api/employees/{a['id']}")["employee"]["designation"] is None  # preview never saves
    api.post("/api/employees/bulk-update", {"updates": upd}, status=422)
    assert api.get(f"/api/employees/{a['id']}")["employee"]["designation"] is None  # atomic: nothing saved
    upd[1]["changes"] = {"designation": "Y"}
    api.post("/api/employees/bulk-update", {"updates": upd})
    assert api.get(f"/api/employees/{b['id']}")["employee"]["designation"] == "Y"
