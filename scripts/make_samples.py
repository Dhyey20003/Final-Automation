"""Generate synthetic sample data: import templates, rosters, placeholder headshots, a ZIP and a PDF.

All people are fictional. Email addresses use the reserved .test domain.
Run:  python scripts/make_samples.py
"""
from __future__ import annotations

import csv
import io
import random
import shutil
import zipfile
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples"
STATIC = ROOT / "hub" / "static" / "samples"

FIRST = ["Aarav", "Diya", "Ishaan", "Kavya", "Rohan", "Ananya", "Vihaan", "Meera", "Arjun", "Saanvi", "Kabir", "Nisha", "Dev",
         "Pooja", "Yash", "Riya", "Neel", "Tara", "Harsh", "Isha", "Manav", "Jiya", "Parth", "Sneha", "Aditya", "Krina", "Om",
         "Hetal", "Rahul", "Zoya", "Kunal", "Bhavya", "Siddharth", "Aditi", "Varun", "Mahi", "Tanmay", "Drashti", "Jay", "Palak"]
LAST = ["Shah", "Patel", "Mehta", "Desai", "Joshi", "Trivedi", "Pandya", "Parikh", "Vyas", "Bhatt", "Modi", "Dave", "Soni", "Rana",
        "Chauhan", "Thakkar", "Gandhi", "Kapadia", "Nair", "Iyer"]
DESIG = ["Associate", "Senior Associate", "Team Lead", "Analyst", "Senior Analyst", "Accountant", "Recruiter", "Developer", "QA Engineer", "Designer"]
DEPT = ["Accounting", "Recruitment", "Engineering", "Operations", "Design", "Finance"]
CLIENTS = ["Northwind Traders", "Contoso Ltd", "Fabrikam Inc", "Tailspin Toys", "Litware Inc"]
HEADERS = ["RHC ID", "Employee Name", "Designation", "Department", "Official Email", "Location", "Date of Joining", "Date of Birth",
           "Status", "Assignment Status", "Client Name", "Assignment Start", "Bench Start", "Remarks"]


def people(n=44, seed=7):
    rnd = random.Random(seed)
    out = []
    used = set()
    for i in range(1, n + 1):
        while True:
            fn, ln = rnd.choice(FIRST), rnd.choice(LAST)
            if (fn, ln) not in used:
                used.add((fn, ln))
                break
        rhc = f"RHC{i:04d}"
        if i == 44:
            rhc = "0457"  # purely numeric ID with a leading zero
        loc = "Ahmedabad" if i % 3 else "Vadodara"
        doj = date(2019, 1, 1) + timedelta(days=rnd.randint(0, 2400))
        dob = date(1985, 1, 1) + timedelta(days=rnd.randint(0, 5000))
        a = rnd.choice(["Client", "Client", "Bench", "Internal"])
        out.append({
            "rhc": rhc, "name": f"{fn} {ln}", "desig": rnd.choice(DESIG), "dept": rnd.choice(DEPT),
            "email": f"{fn.lower()}.{ln.lower()}{i}@relay-sample.test", "loc": loc, "doj": doj, "dob": dob,
            "assign": a, "client": rnd.choice(CLIENTS) if a == "Client" else "",
            "astart": (doj + timedelta(days=30)) if a == "Client" else None,
            "bstart": date(2026, 9, rnd.randint(1, 20)) if a == "Bench" else None,
        })
    return out


def row(p, status="Active", date_fmt="%d-%m-%Y"):
    f = lambda d: d.strftime(date_fmt) if d else ""
    return [p["rhc"], p["name"], p["desig"], p["dept"], p["email"], p["loc"], f(p["doj"]), f(p["dob"]), status, p["assign"],
            p["client"], f(p["astart"]), f(p["bstart"]), ""]


def write_xlsx(path: Path, headers, rows, title="Employees", notes=None, text_cols=(0,)):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B4DB8")
    for r in rows:
        ws.append(r)
    for col in text_cols:
        for cell in ws.iter_cols(min_col=col + 1, max_col=col + 1, min_row=2):
            for c in cell:
                c.number_format = "@"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + i)].width = 18
    if notes:
        ns = wb.create_sheet("Instructions")
        for line in notes:
            ns.append([line])
        ns.column_dimensions["A"].width = 120
    wb.save(path)


def write_csv(path: Path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def headshot(p, size=(400, 500), seed=0) -> Image.Image:
    rnd = random.Random(hash(p["rhc"]) + seed)
    base = (rnd.randint(60, 200), rnd.randint(60, 200), rnd.randint(120, 230))
    img = Image.new("RGB", size, base)
    d = ImageDraw.Draw(img)
    for y in range(size[1]):
        t = y / size[1]
        d.line([(0, y), (size[0], y)], fill=tuple(int(c * (1 - 0.35 * t)) for c in base))
    skin = rnd.choice([(224, 187, 150), (198, 150, 110), (170, 120, 85), (240, 205, 170)])
    d.ellipse([130, 90, 270, 250], fill=skin)
    d.rounded_rectangle([70, 270, 330, 520], radius=90, fill=(40, 50, 80))
    try:
        font = ImageFont.load_default(size=28)
    except TypeError:
        font = ImageFont.load_default()
    initials = "".join(w[0] for w in p["name"].split()[:2])
    d.text((20, 20), initials, fill=(255, 255, 255), font=font)
    d.text((20, size[1] - 40), f"SAMPLE {p['rhc']}", fill=(255, 255, 255), font=font)
    return img


def main():
    OUT.mkdir(exist_ok=True)
    STATIC.mkdir(parents=True, exist_ok=True)
    ppl = people()
    notes = [
        "Relay Employee Hub import template.",
        "RHC ID is required and is kept as text, so leading zeros (e.g. 0457) are preserved.",
        "Dates: DD-MM-YYYY, YYYY-MM-DD or real Excel dates. Ambiguous dates like 03-04-2026 are read day first.",
        "Location: Ahmedabad or Vadodara. Status: Active, Left or Pending Review. Assignment: Client, Bench, Internal or Unknown.",
        "Blank cells never erase existing values unless you choose that on the mapping screen.",
        "Formulas are not run; their saved values are used. Save as .xlsx (not .xlsm or .xls).",
    ]
    write_xlsx(STATIC / "employee_import_template.xlsx", HEADERS, [row(ppl[0])], notes=notes)

    # Historical data: 10 former employees + 15 current employees with older details.
    former = people(10, seed=99)
    for i, p in enumerate(former, start=1):
        p["rhc"] = f"RHC9{i:03d}"
        p["email"] = p["email"].replace("@", f".old{i}@")
    hist_rows = [row(p, "Left") for p in former]
    for p in ppl[:15]:
        r = row(p, "Active")
        r[2] = "Trainee"  # an older designation; current data is kept when importing history after the roster
        hist_rows.append(r)
    write_xlsx(OUT / "synthetic_historical_data.xlsx", HEADERS, hist_rows, title="History", notes=notes)

    # Current complete roster: everyone except the last three (who have left) plus two new joiners.
    current = ppl[:-3]
    roster_rows = [row(p) for p in current]
    joiners = people(2, seed=123)
    for i, p in enumerate(joiners, start=1):
        p["rhc"] = f"RHC{500 + i:04d}"
        p["email"] = p["email"].replace("@", f".new{i}@")
        roster_rows.append(row(p))
    write_csv(OUT / "synthetic_current_roster.csv", HEADERS, roster_rows)
    shutil.copy2(OUT / "synthetic_current_roster.csv", STATIC / "synthetic_current_roster.csv")
    write_csv(OUT / "synthetic_initial_roster.csv", HEADERS, [row(p) for p in ppl])

    # Partial update with blanks (blank must not erase) and a designation update file.
    part = [[ppl[0]["rhc"], "", "Senior Associate", "", "", "", "", "", "", "Bench", "", "", "15-09-2026", ""],
            [ppl[1]["rhc"], "", "", "Finance", "", "", "", "", "", "", "", "", "", "Moved team"],
            ["RHC7777", "Unknown Person", "Analyst", "", "", "Ahmedabad", "", "", "", "", "", "", "", ""]]
    write_csv(OUT / "synthetic_partial_update.csv", HEADERS, part)
    write_csv(OUT / "synthetic_designation_update.csv", ["RHC ID", "New Designation"],
              [[p["rhc"], "Senior " + p["desig"] if not p["desig"].startswith("Senior") else p["desig"]] for p in ppl[:6]])

    # Placeholder headshots: filenames with RHC IDs, names only, and camera-style names.
    photos = OUT / "synthetic_headshots"
    if photos.exists():
        shutil.rmtree(photos)
    photos.mkdir()
    mapping = []
    for i, p in enumerate(ppl[:36]):
        img = headshot(p)
        if i < 24:
            name = f"{p['rhc']}_{p['name'].split()[0]}.jpg"
        elif i < 30:
            name = f"{p['name']}.png"
        else:
            name = f"IMG_{1000 + i}.jpg"
            mapping.append((name, p["rhc"]))
        img.save(photos / name, "PNG" if name.endswith(".png") else "JPEG", quality=90)
    with open(OUT / "synthetic_photo_mapping.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Filename", "RHC ID"])
        w.writerows(mapping)
    with zipfile.ZipFile(OUT / "synthetic_headshots.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for fpath in sorted(photos.iterdir()):
            z.write(fpath, f"headshots/{fpath.name}")

    # PDF with embedded photos and captions, plus a scanned-style page.
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        buf_imgs = []
        c = canvas.Canvas(str(OUT / "synthetic_headshots.pdf"), pagesize=A4)
        W, H = A4
        c.setFont("Helvetica-Bold", 16)
        c.drawString(40, H - 50, "Sample staff photo sheet (synthetic)")
        c.setFont("Helvetica", 10)
        for k, p in enumerate(ppl[36:42]):
            col, rowi = k % 3, k // 3
            x, y = 40 + col * 180, H - 300 - rowi * 260
            tmp = OUT / f"_tmp_{k}.jpg"
            headshot(p).save(tmp, "JPEG")
            buf_imgs.append(tmp)
            c.drawImage(str(tmp), x, y, width=140, height=175)
            c.drawString(x, y - 14, p["name"])
            c.drawString(x, y - 26, p["rhc"])
        c.showPage()
        scan = Image.new("RGB", (1240, 1754), (250, 250, 248))
        for k, p in enumerate(ppl[42:44]):
            scan.paste(headshot(p).resize((300, 375)), (120 + k * 500, 200))
        tmp = OUT / "_tmp_scan.jpg"
        scan.save(tmp, "JPEG")
        buf_imgs.append(tmp)
        c.drawImage(str(tmp), 0, 0, width=W, height=H)
        c.showPage()
        c.save()
        for t in buf_imgs:
            t.unlink()
    except ImportError:
        print("reportlab not installed: skipping the sample PDF")
    print("Samples written to", OUT)


if __name__ == "__main__":
    main()
